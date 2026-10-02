from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "Cisco_ACI_LAB_Guide.md"
OUT = ROOT / "output" / "docx" / "Cisco_ACI_LAB_Guide.docx"
ASSETS = ROOT / "output" / "assets"


def font(size: int, bold: bool = False):
    candidates = [Path("C:/Windows/Fonts/msjhbd.ttc" if bold else "C:/Windows/Fonts/msjh.ttc"), Path("C:/Windows/Fonts/arial.ttf")]
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def box(draw, xy, text, fill, width=2):
    draw.rounded_rectangle(xy, radius=14, fill=fill, outline="#1f2937", width=width)
    x1, y1, x2, y2 = xy
    lines = text.split("\n")
    f = font(28, True)
    total = len(lines) * 38
    for i, line in enumerate(lines):
        bbox = draw.textbbox((0, 0), line, font=f)
        draw.text(((x1 + x2 - (bbox[2] - bbox[0])) / 2, y1 + (y2 - y1 - total) / 2 + i * 38), line, fill="#111827", font=f)


def arrow(draw, a, b, label=""):
    draw.line([a, b], fill="#334155", width=5)
    if label:
        f = font(20)
        bbox = draw.textbbox((0, 0), label, font=f)
        mx, my = (a[0] + b[0]) / 2, (a[1] + b[1]) / 2
        draw.rectangle((mx - 8, my - 18, mx + bbox[2] + 8, my + 14), fill="white")
        draw.text((mx, my - 14), label, fill="#111827", font=f)


def make_diagrams():
    ASSETS.mkdir(parents=True, exist_ok=True)
    im = Image.new("RGB", (1600, 1000), "white")
    d = ImageDraw.Draw(im)
    box(d, (620, 60, 980, 190), "POC-S101\nSpine 101", "#dbeafe")
    box(d, (280, 350, 680, 500), "POC-L201\nLeaf 201", "#dcfce7")
    box(d, (920, 350, 1320, 500), "POC-L202\nLeaf 202", "#dcfce7")
    for i, name in enumerate(["APIC1 .1", "APIC2 .2", "APIC3 .3"]):
        box(d, (40 + i * 230, 680, 230 + i * 230, 790), name, "#fef3c7")
        arrow(d, (135 + i * 230, 680), (420 + i * 40, 500))
    box(d, (720, 680, 1050, 810), "POC-SRV1\n.31", "#f3e8ff")
    box(d, (1170, 680, 1500, 810), "POC-SRV2\n.32", "#f3e8ff")
    arrow(d, (600, 350), (750, 190), "eth1/53")
    arrow(d, (1000, 350), (850, 190), "eth1/53")
    arrow(d, (820, 680), (520, 500), "vmnic2")
    arrow(d, (940, 680), (1050, 500), "vmnic3")
    arrow(d, (1270, 680), (600, 500), "vmnic2")
    arrow(d, (1390, 680), (1150, 500), "vmnic3")
    d.text((50, 900), "Static Port: Leaf 201/202 eth1/1-2   |   eth1/3-4 reserved for VMM", fill="#334155", font=font(28))
    im.save(ASSETS / "physical_topology.png")

    im = Image.new("RGB", (1600, 720), "white")
    d = ImageDraw.Draw(im)
    box(d, (80, 220, 480, 460), "EPG_WEB  VLAN 2101\nBD_WEB  10.1.0.254/24", "#dbeafe")
    box(d, (600, 220, 1000, 460), "EPG_AP  VLAN 2201\nBD_AP  10.2.0.254/24", "#dcfce7")
    box(d, (1120, 220, 1520, 460), "EPG_DB  VLAN 2301\nBD_DB  10.3.0.254/24", "#fef3c7")
    arrow(d, (480, 340), (600, 340), "web_ap")
    arrow(d, (1000, 340), (1120, 340), "ap_db")
    d.text((650, 80), "TN_POC / VRF_POC / AP_POC", fill="#111827", font=font(34, True))
    d.text((430, 600), "All subnets: Public  |  Unicast Routing + ARP Flood + L2 Unknown Flood", fill="#334155", font=font(25))
    im.save(ASSETS / "tenant_topology.png")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side in ("top", "start", "bottom", "end"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), "100")
        node.set(qn("w:type"), "dxa")
        mar.append(node)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    for name, size, bold in [("Normal", 10.5, False), ("Title", 28, True), ("Heading 1", 19, True), ("Heading 2", 14, True)]:
        s = doc.styles[name]
        s.font.name = "Microsoft JhengHei"
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        s.font.size = Pt(size)
        s.font.bold = bold
        s.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Normal"].paragraph_format.space_after = Pt(5)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.15
    doc.styles["Heading 1"].paragraph_format.page_break_before = False
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(8)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(10)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(5)
    add_page_number(section.footer.paragraphs[0])


def add_table(doc, rows):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for i, value in enumerate(rows[0]):
        cell = table.rows[0].cells[i]
        cell.text = value.strip()
        set_cell_shading(cell, "1F4E78")
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True
    for r_idx, row in enumerate(rows[1:]):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value.strip()
            if r_idx % 2:
                set_cell_shading(cells[i], "EAF2F8")
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = "Microsoft JhengHei"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
                    run.font.size = Pt(8.5)
    doc.add_paragraph()


def strip_md(text):
    return re.sub(r"`([^`]+)`", r"\1", re.sub(r"\*\*([^*]+)\*\*", r"\1", text))


def build():
    make_diagrams()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    style_document(doc)
    code = False
    code_lines = []
    i = 0
    mermaid_count = 0
    chapter_count = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            if code:
                p = doc.add_paragraph()
                p.style = doc.styles["Normal"]
                for idx, c in enumerate(code_lines):
                    r = p.add_run(c + ("\n" if idx < len(code_lines) - 1 else ""))
                    r.font.name = "Consolas"
                    r.font.size = Pt(8.5)
                code = False
                code_lines = []
            elif line == "```mermaid":
                i += 1
                while i < len(lines) and not lines[i].startswith("```"):
                    i += 1
                path = ASSETS / ("physical_topology.png" if mermaid_count == 0 else "tenant_topology.png")
                mermaid_count += 1
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(path), width=Inches(6.7))
            else:
                code = True
            i += 1
            continue
        if code:
            code_lines.append(line)
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[-:| ]+\|$", lines[i + 1]):
            rows = [[strip_md(x.strip()) for x in line.strip("|").split("|")]]
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                rows.append([strip_md(x.strip()) for x in lines[i].strip("|").split("|")])
                i += 1
            add_table(doc, rows)
            continue
        if line.startswith("# "):
            text = line[2:]
            if not doc.paragraphs:
                doc.add_heading(text, 0)
                p = doc.add_paragraph("ACI 初學者實體環境操作手冊與章節式自動化工具")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                if chapter_count > 0:
                    doc.add_page_break()
                doc.add_heading(text, 1)
                chapter_count += 1
        elif line.startswith("## "):
            doc.add_heading(line[3:], 2)
        elif line.startswith("### "):
            doc.add_heading(line[4:], 3)
        elif re.match(r"^\d+\. ", line):
            doc.add_paragraph(strip_md(re.sub(r"^\d+\. ", "", line)), style="List Number")
        elif line.startswith("- "):
            doc.add_paragraph(strip_md(line[2:]), style="List Bullet")
        elif line.startswith("> "):
            p = doc.add_paragraph(strip_md(line[2:]))
            p.paragraph_format.left_indent = Inches(0.3)
            for run in p.runs:
                run.font.bold = True
        elif line.strip():
            doc.add_paragraph(strip_md(line))
        i += 1
    doc.core_properties.title = "Cisco ACI LAB Guide"
    doc.core_properties.subject = "Cisco ACI Static Port beginner lab"
    doc.core_properties.author = "ACI LAB"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
